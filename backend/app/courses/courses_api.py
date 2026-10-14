import io
import re
import uuid
import zipfile
from datetime import date
from decimal import Decimal

import openpyxl
from flask import Blueprint, Response, g, request
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, ValidationError
from sqlalchemy.orm import joinedload

from app.auth.auth_utils import require_role, verify_jwt_token
from app.config import FAIL_MESSAGE, SUCCESS_MESSAGE, respond
from app.courses.course_pdf import generate_course_outline_pdf
from app.courses.courses_dao import CourseDAO
from app.courses.courses_import import map_lesson_row, map_row
from app.db import to_dict
from app.models import Course, CourseOutlineStatus, CourseStatus, IPOwner, Instructor, LabLocation, Lesson, SkillLevel
from app.s3_utils import create_presigned_playback_url, create_presigned_upload_post

courses_api = Blueprint("courses_api", __name__, url_prefix="/v1/courses")


class CourseCreateRequest(BaseModel):
    course_name: str
    nfolks_course_code: str
    vendor: str | None = None
    platform: str | None = None
    portfolio: str | None = None
    product: str | None = None
    ibm_course_codes: str | None = None
    status: CourseStatus = CourseStatus.backlog
    course_outline_status: CourseOutlineStatus | None = None
    modality: str | None = None
    duration_days: Decimal | None = None
    num_hours: Decimal | None = None
    on_ibm_pmt: str | None = None
    lab_location: LabLocation | None = None
    ibm_retirement_date: date | None = None
    ip_owner: IPOwner | None = None
    course_author_email: str | None = None
    planned_completion_date: date | None = None
    release_date: date | None = None
    retire_date: date | None = None
    course_description: str | None = None
    link_to_course_outline: str | None = None
    skill_level: SkillLevel | None = None
    max_students: int | None = None
    pre_reqs: str | None = None
    persona_audience_type: str | None = None
    lab_configuration: str | None = None
    course_technical_requirements: str | None = None
    delivery_instruction_notes: str | None = None
    instructor_daily_rate: Decimal | None = None
    materials_cost: Decimal | None = None
    lab_cost: Decimal | None = None
    learning_path: str | None = None
    badge: str | None = None
    certification: str | None = None
    notes: str | None = None
    record_source: str | None = None
    instructor_ids: list[int] | None = None


class CourseUpdateRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")

    course_name: str | None = None
    nfolks_course_code: str | None = None
    vendor: str | None = None
    platform: str | None = None
    portfolio: str | None = None
    product: str | None = None
    ibm_course_codes: str | None = None
    status: CourseStatus | None = None
    course_outline_status: CourseOutlineStatus | None = None
    modality: str | None = None
    duration_days: Decimal | None = None
    num_hours: Decimal | None = None
    on_ibm_pmt: str | None = None
    lab_location: LabLocation | None = None
    ibm_retirement_date: date | None = None
    ip_owner: IPOwner | None = None
    course_author_email: str | None = None
    planned_completion_date: date | None = None
    release_date: date | None = None
    retire_date: date | None = None
    course_description: str | None = None
    link_to_course_outline: str | None = None
    skill_level: SkillLevel | None = None
    max_students: int | None = None
    pre_reqs: str | None = None
    persona_audience_type: str | None = None
    lab_configuration: str | None = None
    course_technical_requirements: str | None = None
    delivery_instruction_notes: str | None = None
    instructor_daily_rate: Decimal | None = None
    materials_cost: Decimal | None = None
    lab_cost: Decimal | None = None
    learning_path: str | None = None
    badge: str | None = None
    certification: str | None = None
    notes: str | None = None
    record_source: str | None = None
    instructor_ids: list[int] | None = None


def course_to_dict(course: Course) -> dict:
    data = to_dict(course)
    instructors = getattr(course, "instructors", None) or []
    data["instructors"] = [to_dict(i) for i in instructors]
    data["instructor_ids"] = [i.id for i in instructors]
    data["instructor"] = ", ".join(i.name for i in instructors if i.name)
    return data


@courses_api.route("", methods=["GET"])
@verify_jwt_token()
def list_courses():
    with CourseDAO() as dao:
        courses = dao.get_all(request.args, g.current_user)
        return respond(SUCCESS_MESSAGE, "OK", [course_to_dict(c) for c in courses])


@courses_api.route("/<int:course_id>", methods=["GET"])
@verify_jwt_token()
def get_course(course_id):
    with CourseDAO() as dao:
        course = dao.get_by_id(course_id, g.current_user)
        if course is None:
            return respond(FAIL_MESSAGE, "Course not found"), 404
        return respond(SUCCESS_MESSAGE, "OK", course_to_dict(course))


@courses_api.route("", methods=["POST"])
@verify_jwt_token()
@require_role("admin", "instructor")
def create_course():
    try:
        body = CourseCreateRequest(**request.get_json(force=True))
    except ValidationError as e:
        return respond(FAIL_MESSAGE, str(e)), 400

    with CourseDAO() as dao:
        course = dao.create(body.model_dump(), g.current_user)
        return respond(SUCCESS_MESSAGE, "Course created", course_to_dict(course)), 201


@courses_api.route("/<int:course_id>", methods=["PUT"])
@verify_jwt_token()
@require_role("admin", "instructor")
def update_course(course_id):
    try:
        body = CourseUpdateRequest(**request.get_json(force=True))
    except ValidationError as e:
        return respond(FAIL_MESSAGE, str(e)), 400

    with CourseDAO() as dao:
        course = dao.get_raw_by_id(course_id)
        if course is None:
            return respond(FAIL_MESSAGE, "Course not found"), 404
        if not dao.owns(course, g.current_user):
            return respond(FAIL_MESSAGE, "You do not own this course"), 403
        updates = body.model_dump(exclude_unset=True)
        if g.current_user["role"] == "instructor":
            updates.pop("course_author_email", None)
        course = dao.update(course, updates)
        return respond(SUCCESS_MESSAGE, "Course updated", course_to_dict(course))


@courses_api.route("/<int:course_id>", methods=["DELETE"])
@verify_jwt_token()
@require_role("admin")
def delete_course(course_id):
    with CourseDAO() as dao:
        course = dao.get_raw_by_id(course_id)
        if course is None:
            return respond(FAIL_MESSAGE, "Course not found"), 404
        dao.delete(course)
        return respond(SUCCESS_MESSAGE, "Course deleted", None)


@courses_api.route("/<int:course_id>/outline", methods=["POST"])
@verify_jwt_token()
@require_role("admin", "instructor")
def request_outline_upload(course_id):
    content_type = request.get_json(force=True).get("content_type", "application/pdf")

    with CourseDAO() as dao:
        course = dao.get_raw_by_id(course_id)
        if course is None:
            return respond(FAIL_MESSAGE, "Course not found"), 404
        if not dao.owns(course, g.current_user):
            return respond(FAIL_MESSAGE, "You do not own this course"), 403

        s3_key = f"outlines/{course_id}/{uuid.uuid4().hex}.pdf"
        upload = create_presigned_upload_post(s3_key, content_type)
        course.course_outline_s3_key = s3_key
        return respond(SUCCESS_MESSAGE, "Upload URL created", {"s3_key": s3_key, "upload": upload})


@courses_api.route("/<int:course_id>/outline", methods=["GET"])
@verify_jwt_token()
def get_outline_url(course_id):
    with CourseDAO() as dao:
        course = dao.get_by_id(course_id, g.current_user)
        if course is None or not course.course_outline_s3_key:
            return respond(FAIL_MESSAGE, "No outline uploaded for this course"), 404
        url = create_presigned_playback_url(course.course_outline_s3_key)
        return respond(SUCCESS_MESSAGE, "OK", {"playback_url": url})


@courses_api.route("/<int:course_id>/outline/pdf", methods=["GET"])
@verify_jwt_token()
def download_course_outline_pdf(course_id):
    with CourseDAO() as dao:
        course = dao.get_by_id(course_id, g.current_user)
        if course is None:
            return respond(FAIL_MESSAGE, "Course not found"), 404
        pdf_bytes = generate_course_outline_pdf(course)
        code = course.nfolks_course_code or f"COURSE-{course.id}"
        safe_code = "".join(c for c in code if c.isalnum() or c in ("-", "_"))
        return Response(
            pdf_bytes,
            mimetype="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{safe_code}_Course_Outline.pdf"',
                "Content-Type": "application/pdf",
            },
        )


@courses_api.route("/import", methods=["POST"])
@verify_jwt_token()
@require_role("admin")
def import_courses():
    file = request.files.get("file")
    if file is None:
        return respond(FAIL_MESSAGE, "No file uploaded (expected multipart field 'file')"), 400

    file_bytes = file.read()
    if not file_bytes:
        return respond(FAIL_MESSAGE, "Uploaded file is empty"), 400

    # 1. Fast XML extraction of cell hyperlinks from .xlsx package
    hyperlinks_map = {}
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as z:
            rels = {}
            for name in z.namelist():
                if name.endswith(".rels") and "sheet" in name:
                    content = z.read(name).decode("utf-8", errors="ignore")
                    for m in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', content):
                        rels[m.group(1)] = m.group(2)
                    for m in re.finditer(r'Target="([^"]+)"[^>]*Id="([^"]+)"', content):
                        rels[m.group(2)] = m.group(1)

            for name in z.namelist():
                if "worksheets/sheet" in name and not name.endswith(".rels"):
                    content = z.read(name).decode("utf-8", errors="ignore")
                    for m in re.finditer(
                        r'<hyperlink ref="([A-Z0-9]+)"(?: [^>]*r:id="([^"]+)")?(?: [^>]*display="([^"]+)")?',
                        content,
                    ):
                        ref = m.group(1)
                        r_id = m.group(2)
                        disp = m.group(3)
                        url = disp or rels.get(r_id)
                        if url:
                            hyperlinks_map[ref] = url
    except Exception:
        pass

    # 2. Fast read_only streaming workbook load
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    # 3. Identify Courses sheet
    courses_sheet_name = None
    for name in workbook.sheetnames:
        clean = name.strip().lower()
        if "course directory" in clean or "course catalog" in clean or clean in ("courses", "course"):
            courses_sheet_name = name
            break
    if not courses_sheet_name:
        courses_sheet_name = workbook.sheetnames[0]

    sheet = workbook[courses_sheet_name]
    row_iter = sheet.iter_rows(values_only=True)
    raw_headers = next(row_iter, None)
    if not raw_headers or not any(raw_headers):
        return respond(FAIL_MESSAGE, "Courses sheet has no headers or is empty"), 400

    headers = [str(h).strip() if h is not None else None for h in raw_headers]

    link_col_idx = None
    for idx, h in enumerate(headers):
        if h and str(h).strip().lower() in ("link to course outline", "link_to_course_outline", "outline link"):
            link_col_idx = idx
            break

    created, updated, skipped, errors = 0, 0, 0, []
    lessons_created, lesson_errors = 0, []

    with CourseDAO() as dao:
        seen_codes = set()
        course_code_to_id = {}
        course_name_to_id = {}

        # Populate lookups for existing courses in DB in a single fast query with instructors
        all_courses = dao.session.query(Course).options(joinedload(Course.instructors)).all()
        course_by_id = {c.id: c for c in all_courses}
        existing_codes = {c.nfolks_course_code for c in all_courses if c.nfolks_course_code}
        for c in all_courses:
            if c.nfolks_course_code:
                course_code_to_id[str(c.nfolks_course_code).strip().lower()] = c.id
            if c.course_name:
                course_name_to_id[str(c.course_name).strip().lower()] = c.id

        # Build instructor lookup by name and email
        all_instructors = dao.session.query(Instructor).all()
        instructor_by_id = {inst.id: inst for inst in all_instructors}
        instructor_by_key = {}
        for inst in all_instructors:
            if inst.name:
                instructor_by_key[inst.name.strip().lower()] = inst.id
            if inst.email:
                instructor_by_key[inst.email.strip().lower()] = inst.id

        # Preload lessons for zero-latency in-memory deduplication
        all_lessons = dao.session.query(Lesson).all()
        lesson_by_key = {(l.course_id, str(l.title).strip().lower()): l for l in all_lessons if l.title}

        # Detect single-sheet mode if sheet headers contain lesson attributes
        clean_headers_set = {str(h).strip().lower() for h in headers if h}
        single_sheet_mode = bool(clean_headers_set & {"lesson title", "title", "class name", "video link", "video url", "attachment url"})

        empty_streak = 0
        for row_number, raw_cells in enumerate(row_iter, start=2):
            if not raw_cells or not any(c is not None and str(c).strip() != "" for c in raw_cells):
                empty_streak += 1
                if empty_streak >= 30:
                    break
                continue

            row = {}
            for idx, (h, val) in enumerate(zip(headers, raw_cells)):
                if not h:
                    continue
                if idx == link_col_idx:
                    coord = f"{get_column_letter(idx + 1)}{row_number}"
                    if coord in hyperlinks_map:
                        val = hyperlinks_map[coord]
                row[h] = val

            fields = map_row(row)

            # Check if this row actually contains course identity data (prevents ghost rows from Excel formulas)
            code = fields.get("nfolks_course_code")
            name = fields.get("course_name")
            row_id = fields.get("id")
            if not (code or name or row_id):
                empty_streak += 1
                if empty_streak >= 30:
                    break
                continue

            empty_streak = 0
            raw_instructor = fields.pop("instructor", None)
            resolved_instructor_ids = []
            if raw_instructor:
                parts = [p.strip() for p in str(raw_instructor).split(",") if p.strip()]
                for p in parts:
                    match_id = instructor_by_key.get(p.lower())
                    if not match_id:
                        for key, i_id in instructor_by_key.items():
                            if p.lower() in key or key in p.lower():
                                match_id = i_id
                                break
                    if match_id and match_id not in resolved_instructor_ids:
                        resolved_instructor_ids.append(match_id)

            author_email = fields.get("course_author_email")
            if author_email:
                match_id = instructor_by_key.get(str(author_email).strip().lower())
                if match_id and match_id not in resolved_instructor_ids:
                    resolved_instructor_ids.append(match_id)

            if resolved_instructor_ids:
                fields["instructor_ids"] = resolved_instructor_ids

            code = fields.get("nfolks_course_code")
            name = fields.get("course_name")
            row_id = fields.get("id")

            target_course_id = None
            if row_id and row_id in course_by_id:
                target_course_id = row_id
            elif code and str(code).strip().lower() in course_code_to_id:
                target_course_id = course_code_to_id[str(code).strip().lower()]
            elif not code and name and str(name).strip().lower() in course_name_to_id:
                target_course_id = course_name_to_id[str(name).strip().lower()]

            if target_course_id:
                existing_course = course_by_id.get(target_course_id)
                if existing_course:
                    try:
                        update_payload = dict(fields)
                        update_payload.pop("id", None)
                        # Do not mutate the unique course code of an existing course during bulk update
                        update_payload.pop("nfolks_course_code", None)

                        body = CourseUpdateRequest(**update_payload)
                        clean_update_data = body.model_dump(exclude_unset=True)

                        # Detect actual field modifications: only update if fields changed!
                        changed_fields = {}
                        for k, new_val in clean_update_data.items():
                            if k == "instructor_ids":
                                curr_inst_ids = sorted([i.id for i in existing_course.instructors])
                                if sorted(new_val) != curr_inst_ids:
                                    changed_fields["instructor_ids"] = new_val
                                continue
                            curr_val = getattr(existing_course, k, None)

                            # Normalize Enum to string value for comparison
                            curr_cmp = curr_val.value if hasattr(curr_val, "value") else curr_val
                            new_cmp = new_val.value if hasattr(new_val, "value") else new_val

                            if isinstance(curr_cmp, (int, float, Decimal)) and isinstance(new_cmp, (int, float, Decimal)):
                                try:
                                    if Decimal(str(curr_cmp)) == Decimal(str(new_cmp)):
                                        continue
                                except Exception:
                                    pass
                            elif isinstance(curr_cmp, str) and isinstance(new_cmp, str):
                                if curr_cmp.strip() == new_cmp.strip():
                                    continue
                            elif curr_cmp == new_cmp:
                                continue

                            changed_fields[k] = new_val

                        # Only update and count if changes actually exist
                        if changed_fields:
                            if "instructor_ids" in changed_fields:
                                inst_ids = changed_fields.pop("instructor_ids")
                                existing_course.instructors = [instructor_by_id[i_id] for i_id in inst_ids if i_id in instructor_by_id]
                            for k, v in changed_fields.items():
                                setattr(existing_course, k, v)
                            updated += 1
                    except Exception as e:
                        errors.append({"row": row_number, "reason": f"Update failed: {str(e)}"})
                        skipped += 1
            else:
                if not code and name:
                    code = f"ATLAS-BK-{row_number:03d}"
                    fields["nfolks_course_code"] = code

                if code:
                    try:
                        create_payload = dict(fields)
                        create_payload.pop("id", None)
                        body = CourseCreateRequest(**create_payload)
                        new_course = dao.create(body.model_dump(), g.current_user)
                        code_key = str(new_course.nfolks_course_code).strip().lower()
                        course_code_to_id[code_key] = new_course.id
                        course_by_id[new_course.id] = new_course
                        if new_course.course_name:
                            course_name_to_id[str(new_course.course_name).strip().lower()] = new_course.id
                        target_course_id = new_course.id
                        seen_codes.add(code_key)
                        created += 1
                    except ValidationError as e:
                        errors.append({"row": row_number, "reason": str(e)})
                        skipped += 1
                        continue
                    except Exception as e:
                        errors.append({"row": row_number, "reason": str(e)})
                        skipped += 1
                        continue

            # In single-sheet mode, process lesson columns on this row if present
            if single_sheet_mode and target_course_id:
                l_fields = map_lesson_row(row)
                if l_fields.get("title") or l_fields.get("attachment_url") or l_fields.get("description"):
                    title = l_fields.get("title") or f"Lesson {l_fields.get('sequence_order', 1)}"
                    url = l_fields.get("attachment_url")
                    c_type = l_fields.get("attachment_content_type")
                    if not c_type and url:
                        c_type = "video" if ("youtube" in url or "youtu.be" in url or ".mp4" in url) else "document"
                    elif not c_type:
                        c_type = "article"

                    l_key = (target_course_id, title.strip().lower())
                    existing_lesson = lesson_by_key.get(l_key)
                    if existing_lesson:
                        if l_fields.get("description"):
                            existing_lesson.description = l_fields.get("description")
                        if l_fields.get("duration_minutes"):
                            existing_lesson.duration_minutes = l_fields.get("duration_minutes")
                        if url:
                            existing_lesson.attachment_url = url
                            existing_lesson.attachment_content_type = c_type
                    else:
                        lesson = Lesson(
                            course_id=target_course_id,
                            title=title,
                            description=l_fields.get("description"),
                            sequence_order=l_fields.get("sequence_order", 1),
                            duration_minutes=l_fields.get("duration_minutes"),
                            attachment_url=url,
                            attachment_content_type=c_type,
                        )
                        dao.session.add(lesson)
                        lesson_by_key[l_key] = lesson
                        lessons_created += 1

        # 2. Check for separate Lessons / Classes sheet (if present and different from courses sheet)
        lessons_sheet_name = None
        for name in workbook.sheetnames:
            if name.strip().lower() in ("lessons", "lesson", "classes", "class", "course_lessons", "course lessons"):
                lessons_sheet_name = name
                break

        if lessons_sheet_name and lessons_sheet_name != courses_sheet_name:
            lesson_sheet = workbook[lessons_sheet_name]
            lesson_rows = lesson_sheet.iter_rows(values_only=True)
            try:
                l_headers = next(lesson_rows)
            except StopIteration:
                l_headers = None

            l_empty_streak = 0
            if l_headers:
                for l_row_num, raw_l_row in enumerate(lesson_rows, start=2):
                    if raw_l_row is None or all(c is None or str(c).strip() == "" for c in raw_l_row):
                        l_empty_streak += 1
                        if l_empty_streak >= 30:
                            break
                        continue
                    l_dict = dict(zip(l_headers, raw_l_row))
                    l_fields = map_lesson_row(l_dict)

                    code = l_fields.get("course_code")
                    if not code:
                        l_empty_streak += 1
                        if l_empty_streak >= 30:
                            break
                        lesson_errors.append({"row": l_row_num, "reason": "Missing course code for lesson"})
                        continue

                    l_empty_streak = 0
                    code_key = str(code).strip().lower()
                    target_course_id = course_code_to_id.get(code_key) or course_name_to_id.get(code_key)
                    if not target_course_id:
                        lesson_errors.append({"row": l_row_num, "reason": f"Course code not found: {code}"})
                        continue

                    title = l_fields.get("title") or f"Lesson {l_fields.get('sequence_order', 1)}"
                    url = l_fields.get("attachment_url")
                    c_type = l_fields.get("attachment_content_type")
                    if not c_type and url:
                        c_type = "video" if ("youtube" in url or "youtu.be" in url or ".mp4" in url) else "document"
                    elif not c_type:
                        c_type = "article"

                    # Deduplicate: check if a lesson with this title already exists for this course
                    l_key = (target_course_id, title.strip().lower())
                    existing_lesson = lesson_by_key.get(l_key)
                    if existing_lesson:
                        if l_fields.get("description"):
                            existing_lesson.description = l_fields.get("description")
                        if l_fields.get("duration_minutes"):
                            existing_lesson.duration_minutes = l_fields.get("duration_minutes")
                        if url:
                            existing_lesson.attachment_url = url
                            existing_lesson.attachment_content_type = c_type
                        continue

                    lesson = Lesson(
                        course_id=target_course_id,
                        title=title,
                        description=l_fields.get("description"),
                        sequence_order=l_fields.get("sequence_order", 1),
                        duration_minutes=l_fields.get("duration_minutes"),
                        attachment_url=url,
                        attachment_content_type=c_type,
                    )
                    dao.session.add(lesson)
                    lesson_by_key[l_key] = lesson
                    lessons_created += 1

        dao.session.commit()

    return respond(
        SUCCESS_MESSAGE,
        "Import complete",
        {
            "created": created,
            "updated": updated,
            "lessons_created": lessons_created,
            "skipped": skipped,
            "errors": errors + lesson_errors,
        },
    )
