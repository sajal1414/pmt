import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_sample_bulk_import_workbook(file_path):
    wb = openpyxl.Workbook()
    
    # Define styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    border_thin = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # ----------------------------------------------------
    # SHEET 1: Courses
    # ----------------------------------------------------
    ws_courses = wb.active
    ws_courses.title = "Courses"
    
    course_headers = [
        "NFOLKS Course Code",
        "Course Name",
        "Vendor",
        "Platform",
        "Portfolio",
        "Product",
        "IBM Course Codes",
        "Modality",
        "Duration (Days)",
        "Duration (Hours)",
        "Skill Level",
        "Max Students",
        "Pre-reqs",
        "Course Description",
        "Course Author Email",
        "Status",
        "Course Outline Status",
        "Lab Location",
        "IP Owner",
        "Planned Completion Date",
        "Release Date",
        "Retire Date",
        "Audience Persona",
        "Lab Configuration",
        "Technical Requirements",
        "Delivery Notes",
        "Daily Rate",
        "Materials Cost",
        "Lab Cost",
        "Learning Path",
        "Badge",
        "Certification",
        "Notes"
    ]
    
    ws_courses.append(course_headers)
    
    courses_data = [
        [
            "ATLAS-CLOUD-301",
            "Advanced Multi-Cloud Architecture & Resilience",
            "NFOLKS",
            "AWS / Azure / GCP",
            "Cloud Infrastructure",
            "Enterprise Hybrid Cloud",
            "IBM-MC-301",
            "Virtual Classroom",
            3,
            24,
            "advanced",
            25,
            "Solid understanding of networking, Kubernetes, and at least one public cloud provider.",
            "Comprehensive deep-dive into multi-region high availability, cross-cloud failover, zero-downtime migrations, and enterprise cloud landing zones.",
            "instructor@nfolks.com",
            "published",
            "validated",
            "nFolks",
            "nFolks",
            "2026-05-15",
            "2026-06-01",
            "2028-12-31",
            "Enterprise Architects, Lead Systems Engineers, Principal Cloud DevOps",
            "Pre-provisioned multi-cloud tenant accounts with Terraform state backends.",
            "Dual-monitor setup recommended, high-speed broadband, SSH client installed.",
            "Hands-on labs require executing Terraform and validating DNS failover routes.",
            1400.00,
            120.00,
            85.00,
            "Cloud Architect Professional Track",
            "NFOLKS Certified Multi-Cloud Architect",
            "NCMA-2026",
            "Includes production blueprints for banking-grade disaster recovery."
        ],
        [
            "ATLAS-DEVOPS-201",
            "Kubernetes GitOps & Production SRE Practices",
            "NFOLKS",
            "Kubernetes & ArgoCD",
            "DevOps & Platform Engineering",
            "ArgoCD / Helm / Prometheus",
            "IBM-GITOPS-201",
            "Hands-on Lab",
            4,
            32,
            "intermediate",
            30,
            "Linux fundamentals, containerization (Docker) and basic Kubernetes knowledge.",
            "Master modern GitOps continuous delivery workflows using ArgoCD, Helm, cert-manager, sealed secrets, and automated canary deployments.",
            "admin@nfolks.com",
            "published",
            "validated",
            "nFolks",
            "nFolks",
            "2026-04-10",
            "2026-04-25",
            "2029-01-01",
            "DevOps Engineers, Site Reliability Engineers, Platform Architects",
            "Dedicated 3-node Kubernetes cluster with ArgoCD pre-installed per participant.",
            "kubectl, helm, and git CLI installed on local workstation.",
            "Includes progressive delivery rollbacks using Argo Rollouts analysis runs.",
            1250.00,
            90.00,
            75.00,
            "Platform Engineering Track",
            "GitOps Master Specialist",
            "GOMS-2026",
            "Designed in collaboration with production SRE team leads."
        ],
        [
            "ATLAS-ML-401",
            "Distributed LLM Training & Real-time Serving",
            "NFOLKS",
            "PyTorch / vLLM / Triton",
            "Artificial Intelligence",
            "DeepSpeed / vLLM",
            "IBM-LLM-401",
            "Virtual Classroom",
            5,
            40,
            "advanced",
            20,
            "Strong Python proficiency, PyTorch experience, deep learning basics.",
            "End-to-end framework for fine-tuning, quantifying, and serving open-source Large Language Models at production scale with vLLM and TensorRT-LLM.",
            "instructor@nfolks.com",
            "published",
            "validated",
            "nFolks",
            "nFolks",
            "2026-07-01",
            "2026-07-15",
            "2028-12-31",
            "AI Engineers, Machine Learning Researchers, Data Platform Leads",
            "GPU instances with 4x NVIDIA A100/H100 GPUs and high-throughput NVMe scratch storage.",
            "Web browser capable of running JupyterLab and modern terminal emulation.",
            "Students will fine-tune a 8B model and benchmark KV cache latency.",
            1800.00,
            250.00,
            320.00,
            "Enterprise GenAI Track",
            "Distributed AI & LLMOps Specialist",
            "DAILLM-2026",
            "Includes enterprise guardrails, speculative decoding, and quantization recipes."
        ],
        [
            "ATLAS-CYBER-302",
            "Zero Trust Network Architecture & SecOps",
            "NFOLKS",
            "Palo Alto / CrowdStrike / Okta",
            "Cybersecurity",
            "Zero Trust Suite",
            "IBM-ZT-302",
            "Hands-on Lab",
            3,
            24,
            "intermediate",
            25,
            "Fundamental understanding of TCP/IP, Active Directory/SAML, and firewall principles.",
            "Practical implementation of Zero Trust architecture across identity (IDP), endpoint posture (EDR), micro-segmentation, and continuous compliance monitoring.",
            "instructor@nfolks.com",
            "published",
            "validated",
            "nFolks",
            "nFolks",
            "2026-08-01",
            "2026-08-10",
            "2029-06-30",
            "SOC Analysts, Cybersecurity Specialists, Infrastructure Security Engineers",
            "Isolated cyber sandbox with simulated attack vectors and active telemetry feeds.",
            "SSH and VPN client compatible with WireGuard / OpenVPN.",
            "Live-fire threat hunting exercises included in day 3.",
            1350.00,
            110.00,
            95.00,
            "Cyber Defense & Security Architecture",
            "Zero Trust Certified Engineer",
            "ZTCE-2026",
            "Covers NIST SP 800-207 Zero Trust framework compliance in detail."
        ]
    ]
    
    for row in courses_data:
        ws_courses.append(row)
        
    # ----------------------------------------------------
    # SHEET 2: Lessons
    # ----------------------------------------------------
    ws_lessons = wb.create_sheet(title="Lessons")
    
    lesson_headers = [
        "Course Code",
        "Sequence Order",
        "Lesson Title",
        "Duration (minutes)",
        "Attachment URL",
        "Content Type",
        "Description"
    ]
    
    ws_lessons.append(lesson_headers)
    
    lessons_data = [
        # ATLAS-CLOUD-301 Lessons
        [
            "ATLAS-CLOUD-301",
            1,
            "Multi-Region Infrastructure Blueprinting",
            45,
            "https://www.youtube.com/watch?v=Ia-UEYYR44s",
            "video",
            "### Architectural Overview\n\nIn this foundational session, we examine the principles of multi-region high availability.\n\n- Active-Active vs Active-Passive topologies\n- Latency budgets and SLA requirements\n- Consensus protocols in distributed stores"
        ],
        [
            "ATLAS-CLOUD-301",
            2,
            "Global Traffic Management & DNS Anycast",
            60,
            "https://www.youtube.com/watch?v=27x0XUv4j2U",
            "video",
            "### Global Routing Deep-Dive\n\nExplore latency-based routing, GeoDNS policies, and cloud health check probes to prevent black-holing traffic during partial outages."
        ],
        [
            "ATLAS-CLOUD-301",
            3,
            "Cross-Cloud Replication & Consensus Guide",
            40,
            "https://www.w3.org/WAI/ER/tests/xhtml/testfiles/resources/pdf/dummy.pdf",
            "document",
            "### Technical Specification Document\n\nDetailed reference guide on configuring cross-cloud object store synchronizers, Kafka mirror-makers, and database read-replicas."
        ],
        [
            "ATLAS-CLOUD-301",
            4,
            "Disaster Recovery Drill & Hands-on Lab Walkthrough",
            75,
            "",
            "article",
            "### Live Disaster Recovery Lab Runbook\n\nFollow these step-by-step instructions:\n\n1. **Inject Partition**: Isolate primary AWS us-east-1 region.\n2. **Verify Route 53**: Confirm automated health check failure within 15 seconds.\n3. **Promote Secondary**: Run failover script for Azure East-US read replica.\n4. **Validate Integrity**: Query customer transaction log consistency."
        ],

        # ATLAS-DEVOPS-201 Lessons
        [
            "ATLAS-DEVOPS-201",
            1,
            "GitOps Philosophy & ArgoCD Core Architecture",
            50,
            "https://www.youtube.com/watch?v=MeU5_r9VjIo",
            "video",
            "### Understanding Declarative Continuous Delivery\n\nWhy Git is the single source of truth for your cluster state. How ArgoCD reconciliation loops continuously fight configuration drift."
        ],
        [
            "ATLAS-DEVOPS-201",
            2,
            "Helm Templating & Multi-Environment Overlays",
            55,
            "https://www.youtube.com/watch?v=4omUKrZcqCc",
            "video",
            "### Environment Management\n\nBest practices for organizing Helm charts, Kustomize bases, and environment overlays (dev, staging, prod) without duplicating manifest code."
        ],
        [
            "ATLAS-DEVOPS-201",
            3,
            "Production Hardening & Sealed Secrets Spec",
            35,
            "https://www.w3.org/WAI/ER/tests/xhtml/testfiles/resources/pdf/dummy.pdf",
            "document",
            "### Security Architecture Whitepaper\n\nAsymmetric cryptography in Git: Managing sensitive credentials with Bitnami Sealed Secrets and HashiCorp Vault integrations."
        ],
        [
            "ATLAS-DEVOPS-201",
            4,
            "Argo Rollouts Automated Canary Deployments",
            65,
            "",
            "article",
            "### Canary Deployment Hands-on Lab\n\nConfigure progressive canary releases with automated rollback triggers based on Prometheus HTTP error rate metrics:\n\n```yaml\napiVersion: argoproj.io/v1alpha1\nkind: Rollout\nmetadata:\n  name: payment-service\nspec:\n  strategy:\n    canary:\n      steps:\n      - setWeight: 10\n      - pause: {duration: 5m}\n      - setWeight: 50\n```"
        ],

        # ATLAS-ML-401 Lessons
        [
            "ATLAS-ML-401",
            1,
            "Foundations of Distributed LLM Architectures",
            60,
            "https://www.youtube.com/watch?v=zjkBMFhNj_g",
            "video",
            "### Distributed Parallelism Paradigms\n\nDetailed analysis of Data Parallelism (DDP), Pipeline Parallelism (PP), and Tensor Parallelism (TP) across multi-GPU compute nodes."
        ],
        [
            "ATLAS-ML-401",
            2,
            "Parameter Efficient Fine-Tuning (LoRA & QLoRA)",
            60,
            "https://www.youtube.com/watch?v=t509sv5MT0w",
            "video",
            "### Fine-Tuning Optimization\n\nPractical implementation of Low-Rank Adaptation (LoRA), 4-bit NormalFloat quantization, and gradient checkpointing to reduce memory footprints by 70%."
        ],
        [
            "ATLAS-ML-401",
            3,
            "High-Throughput Inference with vLLM & PagedAttention",
            50,
            "https://www.youtube.com/watch?v=5CypNCG_k58",
            "video",
            "### Serving Engine Architecture\n\nHow PagedAttention eliminates memory fragmentation in KV caches and increases serving throughput up to 24x over standard Hugging Face pipelines."
        ],
        [
            "ATLAS-ML-401",
            4,
            "Production LLM Benchmark Lab & Guardrails",
            80,
            "",
            "article",
            "### Benchmarking & Guardrails Runbook\n\nRun high-concurrency synthetic load tests against the inference cluster:\n\n- Measure Time to First Token (TTFT) and Inter-Token Latency (ITL)\n- Configure NeMo Guardrails to block prompt injection and PII exfiltration\n- Deploy Prometheus metrics exporter for continuous latency monitoring."
        ],

        # ATLAS-CYBER-302 Lessons
        [
            "ATLAS-CYBER-302",
            1,
            "Zero Trust Core Principles & NIST 800-207 Architecture",
            45,
            "https://www.youtube.com/watch?v=F3N9G57K33I",
            "video",
            "### Deconstructing Perimeter Security\n\nNever Trust, Always Verify: Identity-first security, dynamic risk engines, and continuous authentication models."
        ],
        [
            "ATLAS-CYBER-302",
            2,
            "Micro-segmentation & Software-Defined Perimeters",
            50,
            "https://www.youtube.com/watch?v=gY359qZ7Yf4",
            "video",
            "### Network Isolation at Layer 7\n\nImplementing fine-grained micro-segmentation across cloud VPCs and container mesh networks to prevent lateral attacker movement."
        ],
        [
            "ATLAS-CYBER-302",
            3,
            "Zero Trust Security Policy Playbook",
            30,
            "https://www.w3.org/WAI/ER/tests/xhtml/testfiles/resources/pdf/dummy.pdf",
            "document",
            "### Standard Operating Procedure Guide\n\nStandard rulesets for conditional access policies, device health attestation, and automated quarantine triggers."
        ],
        [
            "ATLAS-CYBER-302",
            4,
            "Live Threat Detection & Incident Containment Lab",
            70,
            "",
            "article",
            "### Incident Response Simulation\n\nIn this scenario, a compromised credential attempts unauthorized lateral traversal:\n\n1. **Anomaly Detected**: User logon from unrecognized IP range with abnormal user agent.\n2. **Step-up Challenge**: Step-up MFA failure triggers immediate credential revocation.\n3. **Host Isolation**: EDR agent isolates workstation within 3 seconds.\n4. **Audit Trail**: Review forensic audit log in SIEM."
        ]
    ]
    
    for row in lessons_data:
        ws_lessons.append(row)
        
    # ----------------------------------------------------
    # Format and Style Sheets
    # ----------------------------------------------------
    for ws in [ws_courses, ws_lessons]:
        # Style headers
        for col_num in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            ws.row_dimensions[1].height = 28
            
        # Style data rows and calculate column widths
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 24
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = data_font
                cell.border = cell_border
                if col in (1, 3, 4, 9, 10, 11, 12, 16, 17, 18, 19, 20, 21, 22, 27, 28, 29):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)

    wb.save(file_path)
    print(f"Successfully generated complete multi-sheet workbook at: {file_path}")

if __name__ == "__main__":
    import os
    paths = [
        r"c:\Users\PrajwalLute\DataDios\nfolks-atlas-ui\public\atlas_complete_bulk_import_with_lessons.xlsx",
        r"c:\Users\PrajwalLute\DataDios\nfolks-atlas-be\atlas_complete_bulk_import_with_lessons.xlsx",
        r"c:\Users\PrajwalLute\DataDios\atlas_complete_bulk_import_with_lessons.xlsx"
    ]
    for p in paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        create_sample_bulk_import_workbook(p)
