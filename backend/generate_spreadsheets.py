import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.db import SessionLocal
from app.models import Course, Lesson


def get_styles():
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    return header_fill, header_font, data_font, cell_border, align_center, align_left, align_wrap


COURSE_HEADERS = [
    "Vendor",
    "Platform",
    "Portfolio",
    "Product",
    "Course Code",
    "IBM Course Codes",
    "Status",
    "Course Name",
    "Course Outline Status",
    "Modality",
    "Duration (days)",
    "Number of Hours",
    "on IBM PMT",
    "Lab Location (ITZ , nFolks, Instructor owned, N/A)",
    "IBM retirement date",
    "IP Owner (nFolks, SubK)",
    "Course Author Email",
    "Planned Completion Date",
    "Release Date",
    "Retire Date",
    "Course Description",
    "Link to Course Outline",
    "Skill Level",
    "Max students",
    "Pre Reqs",
    "Persona",
    "Lab Configuration",
    "Delivery Instruction Notes",
    "Instructor Daily rate",
    "Materials Cost",
    "Lab Cost",
    "Learning Path",
    "Badge",
    "Certification",
    "Notes",
    "Record Source",
]

LESSON_HEADERS = [
    "Course Code",
    "Sequence Order",
    "Lesson Title",
    "Duration (minutes)",
    "Attachment URL",
    "Content Type",
    "Description",
]


def format_workbook(wb):
    header_fill, header_font, data_font, cell_border, align_center, align_left, align_wrap = get_styles()
    for ws in wb.worksheets:
        ws.row_dimensions[1].height = 28
        for col_num in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 24
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = data_font
                cell.border = cell_border
                col_header = ws.cell(row=1, column=col).value
                if col_header in [
                    "Course Code",
                    "IBM Course Codes",
                    "Status",
                    "Course Outline Status",
                    "Modality",
                    "Duration (days)",
                    "Number of Hours",
                    "on IBM PMT",
                    "Lab Location (ITZ , nFolks, Instructor owned, N/A)",
                    "IP Owner (nFolks, SubK)",
                    "Planned Completion Date",
                    "Release Date",
                    "Retire Date",
                    "Skill Level",
                    "Max students",
                    "Sequence Order",
                    "Duration (minutes)",
                    "Content Type",
                ]:
                    cell.alignment = align_center
                elif col_header in ["Course Description", "Description", "Pre Reqs", "Lab Configuration", "Delivery Instruction Notes"]:
                    cell.alignment = align_wrap
                else:
                    cell.alignment = align_left

        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                first_line = val.split("\n")[0]
                if len(first_line) > max_len:
                    max_len = len(first_line)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 55)


def build_all_entries_spreadsheet():
    """Workbook 1: All 20 courses and their 52 lessons currently in the system, with NO attachments."""
    s = SessionLocal()
    try:
        courses = s.query(Course).order_by(Course.id.asc()).all()
        wb = openpyxl.Workbook()

        # Sheet 1: Courses
        ws_courses = wb.active
        ws_courses.title = "Courses"
        ws_courses.append(COURSE_HEADERS)

        # Sheet 2: Lessons
        ws_lessons = wb.create_sheet(title="Lessons")
        ws_lessons.append(LESSON_HEADERS)

        for c in courses:
            status_val = c.status.value if hasattr(c.status, "value") else str(c.status or "")
            outline_status_val = (
                c.course_outline_status.value
                if hasattr(c.course_outline_status, "value")
                else str(c.course_outline_status or "")
            )
            lab_loc_val = c.lab_location.value if hasattr(c.lab_location, "value") else str(c.lab_location or "")
            ip_owner_val = c.ip_owner.value if hasattr(c.ip_owner, "value") else str(c.ip_owner or "")
            skill_level_val = c.skill_level.value if hasattr(c.skill_level, "value") else str(c.skill_level or "")

            course_row = [
                c.vendor or "NFOLKS",
                c.platform or "",
                c.portfolio or "",
                c.product or "",
                c.nfolks_course_code or "",
                c.ibm_course_codes or "",
                status_val,
                c.course_name or "",
                outline_status_val,
                c.modality or "Virtual Classroom",
                float(c.duration_days) if c.duration_days is not None else 3.0,
                float(c.num_hours) if c.num_hours is not None else 24.0,
                c.on_ibm_pmt or "No",
                lab_loc_val or "nFolks",
                str(c.ibm_retirement_date) if c.ibm_retirement_date else "",
                ip_owner_val or "nFolks",
                c.course_author_email or "admin@nfolks.com",
                str(c.planned_completion_date) if c.planned_completion_date else "",
                str(c.release_date) if c.release_date else "",
                str(c.retire_date) if c.retire_date else "",
                c.course_description or "",
                c.link_to_course_outline or "",
                skill_level_val or "intermediate",
                c.max_students or 25,
                c.pre_reqs or "",
                c.persona_audience_type or "",
                c.lab_configuration or "",
                c.delivery_instruction_notes or "",
                float(c.instructor_daily_rate) if c.instructor_daily_rate is not None else 1200.0,
                float(c.materials_cost) if c.materials_cost is not None else 100.0,
                float(c.lab_cost) if c.lab_cost is not None else 80.0,
                c.learning_path or "",
                c.badge or "",
                c.certification or "",
                c.notes or "",
                c.record_source or "Enterprise Catalog",
            ]
            ws_courses.append(course_row)

            lessons = s.query(Lesson).filter_by(course_id=c.id).order_by(Lesson.sequence_order.asc()).all()
            if lessons:
                for l in lessons:
                    lesson_row = [
                        c.nfolks_course_code,
                        l.sequence_order or 1,
                        l.title or f"Lesson {l.sequence_order}",
                        l.duration_minutes or 45,
                        "",  # NO ATTACHMENT URL as requested
                        "article",  # Content type article since no attachment
                        l.description or f"### {l.title}\n\nComprehensive lecture notes and hands-on guidance.",
                    ]
                    ws_lessons.append(lesson_row)
            else:
                for seq in (1, 2):
                    lesson_row = [
                        c.nfolks_course_code,
                        seq,
                        f"{c.course_name} - Module {seq}: Foundation & Implementation",
                        50,
                        "",  # NO ATTACHMENT URL
                        "article",
                        f"### Overview for Module {seq}\n\nKey learning objectives, architectural patterns, and lab exercises.",
                    ]
                    ws_lessons.append(lesson_row)

        format_workbook(wb)
        return wb
    finally:
        s.close()


def build_demo_import_spreadsheet():
    """Workbook 2: A clean set of 10 new demo courses with 30 lessons (3 each), with NO attachments, ready for live import demo."""
    wb = openpyxl.Workbook()

    ws_courses = wb.active
    ws_courses.title = "Courses"
    ws_courses.append(COURSE_HEADERS)

    ws_lessons = wb.create_sheet(title="Lessons")
    ws_lessons.append(LESSON_HEADERS)

    demo_courses = [
        {
            "code": "ATLAS-DEMO-101",
            "name": "Enterprise Agentic AI & Autonomous Systems Architecture",
            "vendor": "NFOLKS",
            "platform": "Artificial Intelligence",
            "portfolio": "Agentic Workflows",
            "product": "LangGraph / AutoGen / CrewAI",
            "ibm_code": "IBM-AG-101",
            "status": "published",
            "outline_status": "validated",
            "modality": "Virtual Classroom",
            "days": 3,
            "hours": 24,
            "skill": "advanced",
            "max_students": 25,
            "desc": "Architecting resilient multi-agent collaborative workflows with memory, reflection, tool invocation, and human-in-the-loop governance for enterprise applications.",
            "lessons": [
                ("Agentic Foundations & Multi-Agent Topologies", 50, "### Foundations of Autonomous Agents\n\n- State machines vs cyclic computational graphs\n- Short-term and long-term agent memory\n- Tool use and OpenAPI integration"),
                ("Hierarchical Supervisor & Worker Patterns", 60, "### Hierarchical Orchestration\n\n- Implementing supervisor routers with LangGraph\n- Delegating tasks to domain-specialist agents\n- Structured schema output validation"),
                ("Enterprise Human-in-the-Loop & Safety Guardrails", 45, "### Production Governance\n\n- Breakpoint execution and state rollback\n- Approval mechanisms before destructive tool execution\n- Latency and token consumption telemetry"),
            ],
        },
        {
            "code": "ATLAS-DEMO-102",
            "name": "Cloud-Native FinOps & Kubernetes Cost Optimization",
            "vendor": "NFOLKS",
            "platform": "Cloud & Infrastructure",
            "portfolio": "FinOps & Efficiency",
            "product": "Kubecost / OpenCost / Prometheus",
            "ibm_code": "IBM-FIN-102",
            "status": "published",
            "outline_status": "validated",
            "modality": "Hands-on Lab",
            "days": 2,
            "hours": 16,
            "skill": "intermediate",
            "max_students": 30,
            "desc": "Real-world strategies to measure, allocate, and optimize multi-tenant Kubernetes and cloud spending using unit economics, right-sizing, and spot automation.",
            "lessons": [
                ("Kubernetes Cost Allocation & Showback Models", 45, "### Granular Cost Accounting\n\n- Allocating shared cluster costs by namespace, label, and service\n- Idle capacity distribution and overhead calculation"),
                ("Workload Right-Sizing & VPA/KEDA Autoscaling", 55, "### Dynamic Scaling Economics\n\n- Configuring Vertical Pod Autoscaler with Goldilocks\n- Event-driven autoscaling with KEDA to avoid over-provisioning"),
                ("Automated Spot Instance & Savings Plan Governance", 60, "### Purchasing Optimization\n\n- Spot orchestration with Karpenter on EKS/GKE\n- Graceful node termination handlers and budget anomaly alerts"),
            ],
        },
        {
            "code": "ATLAS-DEMO-103",
            "name": "Modern Data Mesh & Decentralized Governance",
            "vendor": "NFOLKS",
            "platform": "Data & Analytics",
            "portfolio": "Data Architecture",
            "product": "dbt / Apache Iceberg / Great Expectations",
            "ibm_code": "IBM-MESH-103",
            "status": "published",
            "outline_status": "validated",
            "modality": "Virtual Classroom",
            "days": 4,
            "hours": 32,
            "skill": "advanced",
            "max_students": 20,
            "desc": "Transition from monolithic data pipelines to decentralized domain-driven data products with automated data contracts, federated governance, and CI/CD data testing.",
            "lessons": [
                ("Domain-Oriented Decentralized Data Ownership", 50, "### Principles of Data Mesh\n\n- Defining data products as autonomous architectural units\n- Shift from pipeline-centric to product-centric thinking"),
                ("Implementing Data Contracts with Protobuf & JSON Schema", 65, "### Enforcing Data Contracts\n\n- Declaring schema contracts between producer and consumer services\n- Automated breaking change detection in GitHub Actions"),
                ("Federated Computational Governance & Quality Gates", 60, "### Continuous Quality Assurance\n\n- Great Expectations assertions embedded into dbt transformation jobs\n- Publishing metadata to Apache Atlas / OpenMetadata"),
            ],
        },
        {
            "code": "ATLAS-DEMO-104",
            "name": "Securing Generative AI: OWASP Top 10 for LLMs & Red Teaming",
            "vendor": "NFOLKS",
            "platform": "Cybersecurity",
            "portfolio": "AI Security",
            "product": "LlamaGuard / NeMo Guardrails / Garak",
            "ibm_code": "IBM-SEC-104",
            "status": "published",
            "outline_status": "validated",
            "modality": "Hands-on Lab",
            "days": 3,
            "hours": 24,
            "skill": "advanced",
            "max_students": 20,
            "desc": "Hands-on defense and adversarial red teaming against prompt injection, data exfiltration, indirect poisoning, model inversion, and insecure output handling.",
            "lessons": [
                ("Direct & Indirect Prompt Injection Attacks", 50, "### Adversarial Prompt Engineering\n\n- Jailbreaking techniques and defensive system prompting\n- Indirect injection via retrieval-augmented document parsing"),
                ("Input/Output Guardrails with NeMo & LlamaGuard", 60, "### Programmable AI Firewalls\n\n- Defining topical boundaries and safety moderation filters\n- Sanitizing sensitive PII/tokens before sending to LLM APIs"),
                ("Automated Red Teaming & Continuous Vulnerability Scanning", 70, "### Live Red-Teaming Lab\n\n- Executing automated exploit suites using Garak\n- Benchmarking model hallucination and security regression rates"),
            ],
        },
        {
            "code": "ATLAS-DEMO-105",
            "name": "High-Throughput Streaming with Apache Kafka & Flink",
            "vendor": "NFOLKS",
            "platform": "Data & Analytics",
            "portfolio": "Event-Driven Systems",
            "product": "Kafka / Apache Flink / Schema Registry",
            "ibm_code": "IBM-STREAM-105",
            "status": "published",
            "outline_status": "validated",
            "modality": "Hands-on Lab",
            "days": 4,
            "hours": 32,
            "skill": "intermediate",
            "max_students": 25,
            "desc": "Building sub-second event streaming analytics pipelines with stateful stream processing, exactly-once semantics, windowed aggregations, and schema evolution.",
            "lessons": [
                ("Kafka Cluster Topologies & Partition Rebalancing", 50, "### Distributed Log Internals\n\n- Leader election, ISR lists, and Zero-Copy OS disk writes\n- Partition assignment strategies and consumer group elasticity"),
                ("Stateful Stream Processing with Apache Flink", 65, "### State Management & Checkpointing\n\n- RocksDB state backends and asynchronous barrier snapshotting\n- Handling late-arriving events with watermarks and allowed lateness"),
                ("Exactly-Once End-to-End Processing Guarantees", 60, "### Two-Phase Commit Protocols\n\n- Kafka transactional producers and idempotent sinks\n- End-to-end data integrity validation across stream boundaries"),
            ],
        },
        {
            "code": "ATLAS-DEMO-106",
            "name": "Full-Stack Vector Search & Hybrid RAG Engineering",
            "vendor": "NFOLKS",
            "platform": "Artificial Intelligence",
            "portfolio": "Vector & Hybrid Search",
            "product": "Milvus / Qdrant / BM25 / Cohere Rerank",
            "ibm_code": "IBM-RAG-106",
            "status": "published",
            "outline_status": "validated",
            "modality": "Virtual Classroom",
            "days": 3,
            "hours": 24,
            "skill": "intermediate",
            "max_students": 25,
            "desc": "Designing high-precision retrieval architectures fusing dense vector embeddings, sparse lexical search, reciprocal rank fusion, and cross-encoder rerankers.",
            "lessons": [
                ("Dense & Sparse Embedding Architectures", 45, "### Hybrid Indexing Strategies\n\n- HNSW vs IVF-PQ indexing tradeoffs in vector databases\n- Combining BM25 keyword matching with dense sentence transformers"),
                ("Reciprocal Rank Fusion & Cross-Encoder Reranking", 55, "### Relevance Tuning\n\n- Normalizing and merging disparate score distributions\n- Applying Cohere/BGE rerankers to improve top-K answer retrieval"),
                ("Context Window Management & Chunking Strategies", 60, "### Document Preprocessing\n\n- Semantic chunking vs sliding window approaches\n- Parent-document retrievers and hierarchical summary linking"),
            ],
        },
        {
            "code": "ATLAS-DEMO-107",
            "name": "Cloud Security Posture Management & CIEM with Wiz",
            "vendor": "NFOLKS",
            "platform": "Cybersecurity",
            "portfolio": "Cloud Security",
            "product": "Wiz / AWS IAM Identity Center / Prisma",
            "ibm_code": "IBM-CSPM-107",
            "status": "published",
            "outline_status": "validated",
            "modality": "Hands-on Lab",
            "days": 3,
            "hours": 24,
            "skill": "advanced",
            "max_students": 20,
            "desc": "Agentless cloud risk graph correlation across misconfigurations, vulnerabilities, exposed secrets, and excessive IAM entitlements in AWS, Azure, and GCP.",
            "lessons": [
                ("Security Graph Architecture & Toxic Combinations", 50, "### Correlating Cloud Risk\n\n- Why isolated alerts fail: identifying reachable vulnerabilities with active admin IAM roles\n- Agentless disk snapshot scanning fundamentals"),
                ("Cloud Infrastructure Entitlement Management (CIEM)", 60, "### Least Privilege Enforcement\n\n- Detecting dormant and overprivileged service accounts\n- Automated policy right-sizing based on CloudTrail access history"),
                ("Automated Remediation Workflows with EventBridge", 55, "### Cloud SecOps Playbooks\n\n- Auto-quarantining compromised compute instances\n- Revoking public S3 bucket policies via AWS Lambda"),
            ],
        },
        {
            "code": "ATLAS-DEMO-108",
            "name": "Distributed Microservices Observability with OpenTelemetry",
            "vendor": "NFOLKS",
            "platform": "DevOps & SRE",
            "portfolio": "Observability",
            "product": "OpenTelemetry / Jaeger / Grafana Mimir",
            "ibm_code": "IBM-OTEL-108",
            "status": "published",
            "outline_status": "validated",
            "modality": "Hands-on Lab",
            "days": 3,
            "hours": 24,
            "skill": "intermediate",
            "max_students": 25,
            "desc": "Standardizing traces, metrics, and logs with vendor-neutral OpenTelemetry collectors, context propagation, tail-based sampling, and Grafana Tempo visualization.",
            "lessons": [
                ("OpenTelemetry Architecture & Collector Pipelines", 50, "### Unified Telemetry Ingestion\n\n- Receivers, processors, and exporters in OTel collector configs\n- Batching, memory limiter, and retry queue tuning"),
                ("Distributed Context Propagation & W3C Trace Context", 60, "### Tracing Asynchronous Systems\n\n- Propagating traceparent headers across HTTP and message queues\n- Instrumenting gRPC services and baggage API usage"),
                ("Tail-Based Sampling & High-Cardinality Metrics", 60, "### Efficient Telemetry Storage\n\n- Retaining 100% of error and slow traces while downsampling routine traffic\n- Metric exemplars linking metrics directly to corresponding trace spans"),
            ],
        },
        {
            "code": "ATLAS-DEMO-109",
            "name": "Database Reliability Engineering & High-Availability PostgreSQL",
            "vendor": "NFOLKS",
            "platform": "Data & Infrastructure",
            "portfolio": "Database Engineering",
            "product": "PostgreSQL 16 / Patroni / PgBouncer",
            "ibm_code": "IBM-PG-109",
            "status": "published",
            "outline_status": "validated",
            "modality": "Hands-on Lab",
            "days": 4,
            "hours": 32,
            "skill": "advanced",
            "max_students": 18,
            "desc": "Operating mission-critical PostgreSQL clusters at scale with automated consensus failover via Patroni, connection pooling, VACUUM tuning, and replication monitoring.",
            "lessons": [
                ("PostgreSQL Storage Engine & MVCC Deep Dive", 60, "### Internal Engine Mechanics\n\n- Heap tuples, HOT updates, and transaction wraparound prevention\n- Autovacuum cost-based throttling and table bloating remediation"),
                ("High Availability with Patroni, Etcd, and HAProxy", 70, "### Automated Cluster Failover\n\n- Raft consensus leader election and split-brain protection\n- Zero-downtime rolling minor version upgrades and standby promotions"),
                ("Connection Pooling Architectures with PgBouncer", 50, "### Scalable Session Management\n\n- Session vs transaction pooling modes\n- Mitigating backend connection exhaustion under bursty microservice workloads"),
            ],
        },
        {
            "code": "ATLAS-DEMO-110",
            "name": "Practical Quantum Computing with Qiskit 1.0",
            "vendor": "NFOLKS",
            "platform": "Quantum & Advanced Compute",
            "portfolio": "Quantum Systems",
            "product": "Qiskit / IBM Quantum Platform",
            "ibm_code": "IBM-QC-110",
            "status": "published",
            "outline_status": "validated",
            "modality": "Virtual Classroom",
            "days": 3,
            "hours": 24,
            "skill": "intermediate",
            "max_students": 20,
            "desc": "Hands-on quantum algorithm development using Qiskit 1.0 primitives: Sampler, Estimator, quantum circuit synthesis, error mitigation, and quantum optimization applications.",
            "lessons": [
                ("Quantum Circuit Basics: Superposition & Entanglement", 55, "### Foundational Quantum Gates\n\n- Single-qubit Pauli gates, Hadamard, and phase operations\n- Bell state synthesis and multi-qubit entanglement verification"),
                ("Qiskit 1.0 Runtime Primitives: Estimator & Sampler", 60, "### Modern Qiskit Workflows\n\n- Measuring expectation values with EstimatorV2\n- Error suppression and Zero-Noise Extrapolation (ZNE) techniques"),
                ("Variational Quantum Algorithms for Optimization (QAOA)", 65, "### Industrial Applications\n\n- Formulating combinatorial problems as Ising Hamiltonians\n- Executing QAOA circuits on IBM Quantum superconducting hardware"),
            ],
        },
    ]

    for c in demo_courses:
        course_row = [
            c["vendor"],
            c["platform"],
            c["portfolio"],
            c["product"],
            c["code"],
            c["ibm_code"],
            c["status"],
            c["name"],
            c["outline_status"],
            c["modality"],
            c["days"],
            c["hours"],
            "Yes",
            "nFolks",
            "2028-12-31",
            "nFolks",
            "admin@nfolks.com",
            "2026-06-01",
            "2026-06-15",
            "2028-12-31",
            c["desc"],
            "",
            c["skill"],
            c["max_students"],
            "Basic enterprise architecture experience.",
            "Solutions Architects, Lead Engineers, Technical Directors",
            "Dedicated cloud sandbox environment pre-configured per student.",
            "Distribute credentials 24h before kickoff.",
            1400.0,
            120.0,
            90.0,
            "Enterprise Excellence Track",
            f"{c['vendor']} Certified Specialist",
            f"NC-{c['code'].split('-')[1]}-2026",
            "Demo catalog course ready for instant bulk onboarding.",
            "Curriculum Engineering Q2",
        ]
        ws_courses.append(course_row)

        for seq, (title, dur, desc) in enumerate(c["lessons"], start=1):
            lesson_row = [
                c["code"],
                seq,
                title,
                dur,
                "",  # NO ATTACHMENT URL
                "article",
                desc,
            ]
            ws_lessons.append(lesson_row)

    format_workbook(wb)
    return wb


def main():
    print("Generating Spreadsheet 1: All system entries (no attachments, 2 sheets: Courses & Lessons)...")
    wb1 = build_all_entries_spreadsheet()

    print("Generating Spreadsheet 2: Fresh demo import catalog (no attachments, 2 sheets: Courses & Lessons)...")
    wb2 = build_demo_import_spreadsheet()

    destinations = [
        r"c:\Users\PrajwalLute\DataDios",
        r"c:\Users\PrajwalLute\DataDios\nfolks-atlas-be",
        r"c:\Users\PrajwalLute\DataDios\nfolks-atlas-ui\public",
    ]

    for dest in destinations:
        os.makedirs(dest, exist_ok=True)
        file1_path = os.path.join(dest, "atlas_all_entries_no_attachments.xlsx")
        file2_path = os.path.join(dest, "atlas_demo_import_no_attachments.xlsx")

        wb1.save(file1_path)
        print(f"Saved: {file1_path}")

        wb2.save(file2_path)
        print(f"Saved: {file2_path}")

    print("\nAll spreadsheets successfully prepared!")


if __name__ == "__main__":
    main()
