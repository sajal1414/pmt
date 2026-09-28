"""Import the updated 150-course inventory from 'atlas_course_inventory_latest.xlsx'
into the database, handling formulas, missing codes for backlog courses, and all 36 columns.
"""
import os
import sys
from decimal import Decimal
from datetime import date, datetime
import openpyxl

# Add backend root to sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.db import SessionLocal
from app.models import Course, CourseStatus
from app.courses.courses_import import map_row

def run_import():
    excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "atlas_course_inventory_latest.xlsx")
    if not os.path.exists(excel_path):
        excel_path = r"C:\Users\PrajwalLute\Downloads\nFolks Learning Services Course Inventory 10.9.26 (1).xlsx"
    
    print(f"Loading workbook from: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    
    target_sheet = None
    for name in wb.sheetnames:
        clean = name.strip().lower()
        if "course directory" in clean or "course catalog" in clean or clean in ("courses", "course"):
            target_sheet = name
            break
    if not target_sheet:
        target_sheet = "New Atlas Course Directory" if "New Atlas Course Directory" in wb.sheetnames else wb.sheetnames[0]
        
    print(f"Reading sheet: '{target_sheet}'")
    sheet = wb[target_sheet]
    
    rows_iter = sheet.iter_rows(values_only=True)
    headers = [c for c in next(rows_iter)]
    
    session = SessionLocal()
    imported_count = 0
    updated_count = 0
    
    try:
        for idx, row in enumerate(rows_iter, start=2):
            if not row or not any(row):
                continue
            row_dict = dict(zip(headers, row))
            mapped = map_row(row_dict)
            
            cname = mapped.get("course_name")
            if not cname or str(cname).strip() == "":
                continue
            cname = str(cname).strip()
            
            # Course Code handling: if missing, assign ATLAS-BK-###
            ccode = mapped.get("nfolks_course_code")
            if not ccode or str(ccode).strip() == "":
                ccode = f"ATLAS-BK-{idx:03d}"
                mapped["nfolks_course_code"] = ccode
            else:
                ccode = str(ccode).strip()
                mapped["nfolks_course_code"] = ccode
                
            # Handle duration_days and num_hours formulas
            dur_days = mapped.get("duration_days")
            if dur_days is not None:
                try:
                    dur_days_dec = Decimal(str(dur_days).strip())
                    mapped["duration_days"] = dur_days_dec
                except Exception:
                    mapped["duration_days"] = None
                    dur_days_dec = None
            else:
                dur_days_dec = None
                
            num_hrs = mapped.get("num_hours")
            if num_hrs is not None:
                num_hrs_str = str(num_hrs).strip()
                if num_hrs_str.startswith("="):
                    # Formula like =AtlasCombined[[#This Row],[Duration (days)]]*8
                    if dur_days_dec:
                        mapped["num_hours"] = dur_days_dec * Decimal(8)
                    else:
                        mapped["num_hours"] = Decimal(8)
                else:
                    try:
                        mapped["num_hours"] = Decimal(num_hrs_str)
                    except Exception:
                        mapped["num_hours"] = (dur_days_dec * Decimal(8)) if dur_days_dec else None
            elif dur_days_dec:
                mapped["num_hours"] = dur_days_dec * Decimal(8)

            # Dates parsing helper
            for date_field in ["planned_completion_date", "release_date", "retire_date", "ibm_retirement_date"]:
                val = mapped.get(date_field)
                if val:
                    if isinstance(val, (datetime, date)):
                        mapped[date_field] = val if isinstance(val, date) else val.date()
                    elif isinstance(val, str):
                        try:
                            # Try ISO or standard formats
                            mapped[date_field] = datetime.fromisoformat(val.strip().split("T")[0]).date()
                        except Exception:
                            mapped[date_field] = None
                            
            # Decimal fields
            for dec_field in ["instructor_daily_rate", "materials_cost", "lab_cost"]:
                val = mapped.get(dec_field)
                if val is not None:
                    try:
                        mapped[dec_field] = Decimal(str(val).replace("$", "").replace(",", "").strip())
                    except Exception:
                        mapped[dec_field] = None
                        
            # Integer fields
            for int_field in ["max_students"]:
                val = mapped.get(int_field)
                if val is not None:
                    try:
                        mapped[int_field] = int(val)
                    except Exception:
                        mapped[int_field] = None

            # Check if course exists
            existing = session.query(Course).filter(Course.nfolks_course_code == ccode).first()
            if existing:
                for k, v in mapped.items():
                    setattr(existing, k, v)
                updated_count += 1
            else:
                new_course = Course(**mapped)
                session.add(new_course)
                imported_count += 1
                
        session.commit()
        total_courses = session.query(Course).count()
        print(f"SUCCESS: Imported {imported_count} new courses, updated {updated_count} existing courses.")
        print(f"Total courses now in database: {total_courses}")
    except Exception as e:
        session.rollback()
        print(f"ERROR during import: {e}")
        raise
    finally:
        session.close()
        wb.close()

if __name__ == "__main__":
    run_import()
