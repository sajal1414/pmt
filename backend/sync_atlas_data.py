"""Sync Atlas data script:
1. Ensure 'phone' column exists on instructors table.
2. Remove 4 sample instructors (Dr. Sarah Jenkins, Marcus Vance, Elena Rostova, Vikram Anand) & users.
3. Remove old courses, lessons, deliveries, feedback, and junction records.
4. Add phone numbers for all instructors.
5. Clean placeholder course codes from courses_can_teach.
6. Import all courses from 'nFolks Learning Services Course Inventory 2.10.26 1.xlsx':
   - Set status to only: available, in_development, retired, pipeline, backlog.
   - Link to right instructors in course_instructors.
   - Populate instructors' courses_can_teach with actual Course Names.
"""

import os
import sys
import re
import openpyxl
from decimal import Decimal

sys.path.insert(0, os.path.dirname(__file__))

from app.db import SessionLocal, engine
from app.models import (
    Course,
    CourseStatus,
    Instructor,
    Lesson,
    Delivery,
    DeliveryFeedback,
    User,
    UserRole,
    TravelPreference,
    course_instructors,
)
from app.courses.courses_import import map_row
from app.auth.auth_utils import hash_password

EXCEL_PATH = r"C:\Users\PrajwalLute\Downloads\nFolks Learning Services Course Inventory 2.10.26 1.xlsx"

# Realistic phone number prefixes by country/location keyword
PHONE_PATTERNS = {
    "united states": "+1 (555) {n:03d}-{s:04d}",
    "usa": "+1 (555) {n:03d}-{s:04d}",
    "austin": "+1 (512) {n:03d}-{s:04d}",
    "london": "+44 20 {n:04d} {s:04d}",
    "uk": "+44 20 {n:04d} {s:04d}",
    "germany": "+49 30 {n:04d}{s:03d}",
    "berlin": "+49 30 {n:04d}{s:03d}",
    "india": "+91 9820{n:1d} {s:05d}",
    "south africa": "+27 21 555 {s:04d}",
    "poland": "+48 22 {n:03d} {s:04d}",
    "italy": "+39 06 {n:04d} {s:03d}",
    "spain": "+34 91 {n:03d} {s:04d}",
    "france": "+33 1 {n:04d} {s:04d}",
    "netherlands": "+31 20 {n:03d} {s:04d}",
    "belgium": "+32 2 {n:03d} {s:04d}",
    "switzerland": "+41 22 {n:03d} {s:04d}",
}

def generate_phone(idx: int, location: str | None) -> str:
    loc = (location or "").lower()
    for key, pattern in PHONE_PATTERNS.items():
        if key in loc:
            return pattern.format(n=(100 + idx % 899), s=(1000 + (idx * 37) % 8999))
    return f"+1 (555) {200 + idx % 700:03d}-{1000 + (idx * 43) % 8999:04d}"

def run_sync():
    # 1. Ensure 'phone' column in postgres
    print("1. Ensuring phone column exists on instructors table...")
    with engine.connect() as conn:
        conn.exec_driver_sql("ALTER TABLE instructors ADD COLUMN IF NOT EXISTS phone VARCHAR(50);")
        conn.commit()
    print("   Phone column checked/added.")

    session = SessionLocal()
    try:
        # 2. Delete 4 sample instructors & their users
        print("2. Removing 4 sample instructors and their accounts...")
        sample_emails = [
            "sarah.jenkins@nfolksdata.com",
            "marcus.vance@nfolksdata.com",
            "elena.rostova@nfolksdata.com",
            "vikram.anand@nfolksdata.com",
        ]
        sample_users = session.query(User).filter(User.email.in_(sample_emails)).all()
        for u in sample_users:
            session.delete(u)
        sample_instructors = session.query(Instructor).filter(Instructor.email.in_(sample_emails)).all()
        for i in sample_instructors:
            session.delete(i)
        session.commit()
        print(f"   Removed sample instructors & user accounts.")

        # 3. Clean courses, lessons, deliveries, feedback
        print("3. Removing old courses, lessons, deliveries, feedback, and links...")
        session.query(DeliveryFeedback).delete()
        session.query(Delivery).delete()
        session.query(Lesson).delete()
        session.execute(course_instructors.delete())
        session.query(Course).delete()
        session.commit()
        print("   Old course tables cleared.")

        # 4. Clean placeholder course codes from all remaining instructors and assign phone numbers
        print("4. Updating instructors: phone numbers & clearing placeholder course codes...")
        instructors = session.query(Instructor).all()
        for idx, inst in enumerate(instructors, start=1):
            if not inst.phone:
                inst.phone = generate_phone(idx, inst.location)
            # Remove placeholder codes like ATLAS-..., HAZ-...
            if inst.courses_can_teach:
                codes = [c.strip() for c in inst.courses_can_teach.split(",") if c.strip()]
                # If these are just codes like ATLAS-AI-101 or HAZ-AT-001, clear them
                clean_codes = [c for c in codes if not re.match(r"^[A-Z]{3,5}-[A-Z0-9]+-\d+$", c)]
                inst.courses_can_teach = ", ".join(clean_codes) if clean_codes else None
        session.commit()
        print(f"   Updated phone numbers for {len(instructors)} instructors.")

        # 5. Build instructor lookup
        inst_by_name = {}
        inst_by_email = {}
        for inst in session.query(Instructor).all():
            if inst.name:
                inst_by_name[inst.name.strip().lower()] = inst
            if inst.email:
                inst_by_email[inst.email.strip().lower()] = inst

        # Alias fixes for known typos / nicknames in the spreadsheet
        alias_map = {
            "mansih ladha": "manish ladha",
            "niavs kothandapani": "nivas kothandapani",
            "arturosantos": "arturo santos",
            "jacek zimordzoki": "jacek zimordozki",
            "dan ani": "dan gabriel ani",
            "dan ani manish ladha": "special_dan_ani_manish",
        }

        # 6. Read and import courses from Excel
        print(f"6. Loading courses from: {EXCEL_PATH}")
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        ws = wb["New Atlas Course Directory"]
        raw_headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        headers = [str(h).strip() if h is not None else None for h in raw_headers]

        # Map of instructor_id -> list of course names taught
        instructor_taught_courses = {}

        courses_created = 0
        default_pwd_hash = hash_password("nFolks")

        def resolve_or_create_instructor(name_raw: str) -> Instructor:
            name_clean = name_raw.strip()
            name_lower = name_clean.lower()
            name_lower = alias_map.get(name_lower, name_lower)

            if name_lower in inst_by_name:
                return inst_by_name[name_lower]
            for k, inst in inst_by_name.items():
                if name_lower in k or k in name_lower:
                    return inst

            # Not found, create new instructor
            slug = re.sub(r"[^a-z0-9]", ".", name_lower).strip(".")
            email = f"{slug}@nfolksdata.com"
            new_phone = generate_phone(len(inst_by_name) + 50, "United States")
            print(f"   Creating new instructor record: {name_clean} <{email}>")
            new_inst = Instructor(
                name=name_clean,
                email=email,
                phone=new_phone,
                location="United States",
                travel_preference=TravelPreference.both,
                vendor="nFolks",
                platform="Enterprise",
            )
            session.add(new_inst)
            session.flush()

            # Create User account
            new_user = User(
                email=email,
                password_hash=default_pwd_hash,
                role=UserRole.instructor,
                full_name=name_clean,
                instructor_id=new_inst.id,
                must_change_password=True,
            )
            session.add(new_user)
            session.flush()

            inst_by_name[name_lower] = new_inst
            inst_by_email[email.lower()] = new_inst
            return new_inst

        seen_codes = set()
        for r in range(2, ws.max_row + 1):
            row = {}
            for idx, h in enumerate(headers, start=1):
                if h:
                    row[h] = ws.cell(r, idx).value

            fields = map_row(row)
            code = fields.get("nfolks_course_code")
            name = fields.get("course_name")
            if not code or not name:
                continue

            if code in seen_codes:
                code = f"{code}-B"
                fields["nfolks_course_code"] = code
            seen_codes.add(code)

            raw_instructor = fields.pop("instructor", None)
            course_instructors_list = []

            if raw_instructor:
                # Handle special case 'Dan Ani Manish Ladha'
                raw_str = str(raw_instructor)
                if "dan ani manish ladha" in raw_str.lower():
                    raw_str = raw_str.replace("Dan Ani Manish Ladha", "Dan Ani; Manish Ladha")
                
                parts = [p.strip() for p in re.split(r"[;,]", raw_str) if p.strip()]
                for p in parts:
                    inst_obj = resolve_or_create_instructor(p)
                    if inst_obj not in course_instructors_list:
                        course_instructors_list.append(inst_obj)

            author_email = fields.get("course_author_email")
            if author_email:
                inst_obj = inst_by_email.get(str(author_email).strip().lower())
                if inst_obj and inst_obj not in course_instructors_list:
                    course_instructors_list.append(inst_obj)

            # Ensure status is strictly one of the 5 allowed statuses
            status_val = fields.get("status", "backlog")
            if hasattr(status_val, "value"):
                status_val = status_val.value
            status_val = str(status_val).lower().replace(" ", "_").replace("-", "_")
            if status_val in ("available", "published", "active", "live"):
                status_enum = CourseStatus.available
            elif status_val in ("in_development", "in_progress", "for_review", "updated_version", "custom", "lockheed", "venu"):
                status_enum = CourseStatus.in_development
            elif status_val in ("retired", "archive", "archived"):
                status_enum = CourseStatus.retired
            elif status_val == "pipeline":
                status_enum = CourseStatus.pipeline
            else:
                status_enum = CourseStatus.backlog

            fields["status"] = status_enum

            # Create course
            course = Course(**fields)
            course.instructors = course_instructors_list
            session.add(course)
            session.flush()

            # Record course for instructor
            for inst_obj in course_instructors_list:
                if inst_obj.id not in instructor_taught_courses:
                    instructor_taught_courses[inst_obj.id] = []
                if name not in instructor_taught_courses[inst_obj.id]:
                    instructor_taught_courses[inst_obj.id].append(name)

            courses_created += 1

        # 7. Update courses_can_teach for all instructors with the real Course Names!
        print("7. Updating instructors' courses_can_teach with actual Course Names...")
        for inst_id, c_names in instructor_taught_courses.items():
            inst = session.query(Instructor).filter(Instructor.id == inst_id).one_or_none()
            if inst:
                existing = [x.strip() for x in (inst.courses_can_teach or "").split(",") if x.strip()]
                merged = list(dict.fromkeys(existing + c_names))
                inst.courses_can_teach = ", ".join(merged)

        session.commit()
        print(f"SUCCESS: Created {courses_created} courses with verified instructors!")
        print(f"Total instructors now in DB: {session.query(Instructor).count()}")
        print(f"Total courses now in DB: {session.query(Course).count()}")

        # Print breakdown of statuses in DB
        from sqlalchemy import func
        status_counts = session.query(Course.status, func.count(Course.id)).group_by(Course.status).all()
        print("\nCourse Status Distribution in DB:")
        for st, cnt in status_counts:
            print(f"  {st}: {cnt}")

    except Exception as e:
        session.rollback()
        print(f"ERROR: {e}")
        import traceback
        traceback.print_exc()
        raise
    finally:
        session.close()

if __name__ == "__main__":
    run_sync()
