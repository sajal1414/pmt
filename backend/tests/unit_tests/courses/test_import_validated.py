"""Assertion-based self-check test for validated course inventory mapping & outline links."""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))))

import openpyxl
from decimal import Decimal
from app.courses.courses_import import map_row, _STATUS_MAP, _SKILL_LEVEL_MAP
from app.courses.courses_api import CourseCreateRequest, CourseUpdateRequest

def test_status_and_skill_normalization():
    assert _STATUS_MAP.get("lockheed") == "lockheed"
    assert _SKILL_LEVEL_MAP.get("foundation") == "foundational"
    assert _SKILL_LEVEL_MAP.get("intermediate") == "intermediate"

    sample_row = {
        "Course Code": "TEST-001",
        "Course Name": "Test Watsonx Course",
        "Status": "Lockheed",
        "Skill Level": "Foundation",
        "Duration (days)": 3,
        "Number of Hours": 24,
        "Link to Course Outline": "https://nfolksdata.sharepoint.com/:w:/s/TrainingManagement/test?e=123",
    }
    mapped = map_row(sample_row)
    assert mapped["nfolks_course_code"] == "TEST-001"
    assert mapped["status"] == "lockheed"
    assert mapped["skill_level"] == "foundational"
    assert mapped["link_to_course_outline"].startswith("https://nfolksdata.sharepoint.com/")

    # Validate against Pydantic schema
    req = CourseCreateRequest(**mapped)
    assert req.nfolks_course_code == "TEST-001"
    assert req.status.value == "lockheed"
    assert req.skill_level.value == "foundational"

def test_validated_spreadsheet_sample():
    excel_path = r"C:\Users\PrajwalLute\Downloads\nFolks Learning Services Course Inventory 17.9.26 VALIDATED.xlsx"
    if not os.path.exists(excel_path):
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["New Atlas Course Directory"]
    headers = [c.value for c in ws[1]]
    link_idx = headers.index("Link to Course Outline") + 1

    checked = 0
    for r in range(2, ws.max_row + 1):
        raw_cells = [ws.cell(row=r, column=c) for c in range(1, len(headers) + 1)]
        if not any(c.value is not None for c in raw_cells):
            continue
        row_dict = {}
        for idx, (h, cell) in enumerate(zip(headers, raw_cells), 1):
            val = cell.hyperlink.target if (idx == link_idx and cell.hyperlink) else cell.value
            row_dict[h] = val

        mapped = map_row(row_dict)
        assert mapped.get("course_name"), f"Row {r} missing course_name"
        assert mapped.get("nfolks_course_code"), f"Row {r} missing nfolks_course_code"
        assert mapped.get("link_to_course_outline"), f"Row {r} missing link_to_course_outline"
        assert mapped["link_to_course_outline"].startswith("https://nfolksdata.sharepoint.com/"), f"Row {r} link not SharePoint: {mapped.get('link_to_course_outline')}"

        req = CourseUpdateRequest(**mapped)
        assert req is not None
        checked += 1

    wb.close()
    assert checked == 69, f"Expected 69 courses, checked {checked}"
    print(f"Validated all {checked} courses successfully with CourseUpdateRequest and SharePoint outline links!")

if __name__ == "__main__":
    test_status_and_skill_normalization()
    test_validated_spreadsheet_sample()
    print("ALL TESTS PASSED!")
