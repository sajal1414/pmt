import os
import sys
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))))

import openpyxl
from app.instructors.instructors_api import InstructorCreateRequest
from app.instructors.instructors_import import (
    TRAVEL_PREFERENCE_MAP,
    clean_instructor_email,
    map_instructor_row,
)
from app.models import TravelPreference


def test_clean_instructor_email():
    # Regular email
    pe, se = clean_instructor_email("test@example.com", "Test User")
    assert pe == "test@example.com"
    assert se is None

    # Multiple comma-separated emails
    pe, se = clean_instructor_email("primary@example.com, alt@example.com", "Test User")
    assert pe == "primary@example.com"
    assert se == "alt@example.com"

    # mailto: prefix
    pe, se = clean_instructor_email("mailto:someone@yahoo.com", "Someone")
    assert pe == "someone@yahoo.com"
    assert se is None

    # Missing email fallback
    pe, se = clean_instructor_email(None, "Annelies Kobus")
    assert pe == "annelies.kobus@nfolksdata.com"
    assert se is None


def test_map_instructor_row():
    row = {
        "Name": "Heather Stevens",
        "Skills/products": "IBM Cognos,",
        "Email ID": "heathersame@gmail.com",
        "Location": "South Africa",
        "Languages They Teach": "African, English",
        "Travel Preference": "Yes",
        "Vendor": "IBM",
        "Platform": "Data & AI",
    }
    mapped = map_instructor_row(row)
    assert mapped["name"] == "Heather Stevens"
    assert mapped["skills_products"] == "IBM Cognos"
    assert mapped["email"] == "heathersame@gmail.com"
    assert mapped["location"] == "South Africa"
    assert mapped["languages"] == "African, English"
    assert mapped["travel_preference"] == TravelPreference.can_travel
    assert mapped["vendor"] == "IBM"
    assert mapped["platform"] == "Data & AI"

    # Validate against Pydantic schema
    req = InstructorCreateRequest(**mapped)
    assert req.name == "Heather Stevens"
    assert req.travel_preference == TravelPreference.can_travel


def test_ibm_excel_file():
    excel_path = r"C:\Users\PrajwalLute\Downloads\IBM New Instructor list for nFolks ATLAS.xlsx"
    if not os.path.exists(excel_path):
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]

    valid_count = 0
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if not any(vals):
            continue
        row_dict = dict(zip(headers, vals))
        if not row_dict.get("Name"):
            continue

        mapped = map_instructor_row(row_dict)
        req = InstructorCreateRequest(**mapped)
        assert req.name
        assert req.email
        valid_count += 1

    assert valid_count == 74, f"Expected 74 instructors, got {valid_count}"


if __name__ == "__main__":
    test_clean_instructor_email()
    test_map_instructor_row()
    test_ibm_excel_file()
    print("All instructor import unit tests passed!")
